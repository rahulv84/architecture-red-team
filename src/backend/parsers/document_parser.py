"""Plain text, Markdown, DOCX and PPTX parsers."""
from __future__ import annotations

import io
import re
import zipfile

from ..core.exceptions import ParseError
from .types import DocSection, ParsedDocument, ParserLimits

_HEADING = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")


def decode_text(data: bytes, limits: ParserLimits) -> str:
    """Decode UTF-8 (with BOM tolerance); reject binary content."""
    if b"\x00" in data[:4096]:
        raise ParseError("File appears to be binary, not UTF-8 text.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseError(f"File is not valid UTF-8 text: {exc}") from exc
    if len(text) > limits.max_text_chars:
        raise ParseError(f"Text exceeds the {limits.max_text_chars} character limit.")
    return text


def sections_from_markdown(text: str) -> list[DocSection]:
    """Split on ATX headings, recording 1-based line ranges."""
    lines = text.splitlines()
    sections: list[DocSection] = []
    heading: str | None = None
    start = 1
    buf: list[str] = []
    in_fence = False

    def flush(end_line: int) -> None:
        body = "\n".join(buf).strip()
        if body or heading:
            sections.append(DocSection(heading=heading, text=body, line_start=start, line_end=max(start, end_line)))

    for idx, line in enumerate(lines, start=1):
        if line.strip().startswith("```"):
            in_fence = not in_fence
        match = None if in_fence else _HEADING.match(line)
        if match:
            flush(idx - 1)
            heading = match.group(2)
            start = idx
            buf = []
        else:
            buf.append(line)
    flush(len(lines))
    return sections


def sections_from_plain_text(text: str, chunk_chars: int = 1800) -> list[DocSection]:
    """Group paragraphs into ~chunk_chars sections with line ranges."""
    lines = text.splitlines()
    sections: list[DocSection] = []
    buf: list[str] = []
    start = 1
    size = 0
    for idx, line in enumerate(lines, start=1):
        if not buf:
            start = idx
        buf.append(line)
        size += len(line) + 1
        if size >= chunk_chars and not line.strip():
            body = "\n".join(buf).strip()
            if body:
                sections.append(DocSection(text=body, line_start=start, line_end=idx))
            buf, size = [], 0
    body = "\n".join(buf).strip()
    if body:
        sections.append(DocSection(text=body, line_start=start, line_end=len(lines)))
    return sections


def parse_text(data: bytes, filename: str, limits: ParserLimits, markdown: bool) -> ParsedDocument:
    text = decode_text(data, limits)
    sections = sections_from_markdown(text) if markdown else sections_from_plain_text(text)
    if markdown and not any(s.heading for s in sections):
        sections = sections_from_plain_text(text)
    return ParsedDocument(filename=filename, kind="markdown" if markdown else "text", text=text, sections=sections)


def _guard_zip(data: bytes, limits: ParserLimits) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ParseError("File is not a valid Office (zip) container.") from exc
    infos = archive.infolist()
    if len(infos) > limits.max_zip_entries:
        raise ParseError("Archive contains too many entries.")
    if sum(i.file_size for i in infos) > limits.max_zip_uncompressed:
        raise ParseError("Archive expands beyond the allowed uncompressed size (possible decompression bomb).")
    return archive


def parse_docx(data: bytes, filename: str, limits: ParserLimits) -> ParsedDocument:
    import docx  # local import: optional heavy dependency

    _guard_zip(data, limits)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises many concrete types for corrupt input
        raise ParseError(f"Could not read DOCX: {exc}") from exc
    sections: list[DocSection] = []
    heading: str | None = None
    buf: list[str] = []
    paragraph_no = 0

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body or heading:
            sections.append(DocSection(heading=heading, text=body, line_start=None))

    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text.strip()
        if not text:
            continue
        paragraph_no += 1
        if style.lower().startswith("heading") or style.lower() == "title":
            flush()
            heading, buf = text, []
        else:
            buf.append(text)
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        buf.append("\n".join(rows))
    flush()
    warnings: list[str] = []
    image_parts = [n for n in _guard_zip(data, limits).namelist() if n.startswith("word/media/")]
    requires_vision = bool(image_parts)
    if image_parts:
        warnings.append(f"{len(image_parts)} embedded image(s) not interpreted; vision analysis is needed to read them.")
    text = "\n\n".join(((f"## {s.heading}\n" if s.heading else "") + s.text) for s in sections)
    return ParsedDocument(
        filename=filename, kind="docx", text=text, sections=sections, warnings=warnings,
        requires_vision=requires_vision, metadata={"embedded_images": len(image_parts)},
    )


def parse_pptx(data: bytes, filename: str, limits: ParserLimits) -> ParsedDocument:
    from pptx import Presentation

    _guard_zip(data, limits)
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Could not read PPTX: {exc}") from exc
    sections: list[DocSection] = []
    pictures = 0
    for number, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        title = None
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                pictures += 1
            if shape.has_text_frame:
                txt = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
                if txt:
                    if title is None and getattr(shape, "is_placeholder", False):
                        title = txt.splitlines()[0]
                    parts.append(txt)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Speaker notes] {notes}")
        sections.append(DocSection(heading=title or f"Slide {number}", text="\n".join(parts), slide=number))
    warnings = []
    if pictures:
        warnings.append(f"{pictures} picture(s) on slides were not interpreted; vision analysis is needed to read them.")
    text = "\n\n".join(f"## Slide {s.slide}: {s.heading}\n{s.text}" for s in sections)
    return ParsedDocument(
        filename=filename, kind="pptx", text=text, sections=sections, page_count=len(sections),
        warnings=warnings, requires_vision=pictures > 0, metadata={"pictures": pictures},
    )
