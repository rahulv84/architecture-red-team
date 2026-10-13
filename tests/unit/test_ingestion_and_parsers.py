"""Upload validation (1), text/PDF parsing (2) and diagram/IaC parsing (3)."""
from __future__ import annotations

import base64
import io
import json
import urllib.parse
import zlib

import pytest

from backend.core.exceptions import FileTooLargeError, ParseError, UnsupportedFileError
from backend.parsers.parser_registry import detect_type, parse_bytes, safe_filename
from backend.parsers.types import ParserLimits

LIMITS = ParserLimits()


# ----------------------------------------------------------------- upload validation
def test_safe_filename_strips_paths_and_unsafe_characters():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("C:\\temp\\evil<>.tf") == "evil_.tf"
    assert safe_filename("") == "upload"


def test_detect_type_accepts_supported_and_rejects_unknown():
    assert detect_type("a.tf", b'resource "x" "y" {}') == "terraform"
    assert detect_type("a.mmd", b"graph TD\nA-->B") == "mermaid"
    with pytest.raises(UnsupportedFileError):
        detect_type("payload.exe", b"MZ\x90\x00")


def test_signature_mismatch_is_rejected():
    with pytest.raises(UnsupportedFileError):
        detect_type("fake.pdf", b"not a pdf at all")
    with pytest.raises(UnsupportedFileError):
        detect_type("fake.png", b"GIF89a....")


def test_ingestion_stores_hash_type_role_and_is_idempotent(container):
    from backend.models.domain import Project

    project = container.tables.projects.put(Project(name="p"))
    data = b"# Title\n\nThe system uses PostgreSQL."
    a1 = container.ingestion.add_document(project.project_id, "../../design.md", data, "PRIMARY_ARCHITECTURE")
    assert a1.filename == "design.md"
    assert len(a1.sha256) == 64 and a1.detected_type == "markdown" and a1.role.value == "PRIMARY_ARCHITECTURE"
    assert a1.stored_path.startswith(str(container.paths.uploads))
    a2 = container.ingestion.add_document(project.project_id, "design.md", data)
    assert a2.artifact_id == a1.artifact_id  # same bytes -> same artifact
    changed = container.ingestion.add_document(project.project_id, "design.md", data + b"\nmore", None)
    assert changed.artifact_id != a1.artifact_id
    assert container.tables.artifacts.get(a1.artifact_id).active is False  # superseded, not deleted
    assert any(m["sha256"] == a1.sha256 for m in container.ingestion.manifest(project.project_id))


def test_oversize_and_empty_files_rejected(container):
    from backend.models.domain import Project

    project = container.tables.projects.put(Project(name="p"))
    container.settings_store.update({"max_review_rounds": 3})
    big = b"a" * (container.settings_store.get().max_upload_mb * 1024 * 1024 + 1)
    with pytest.raises(FileTooLargeError):
        container.ingestion.add_document(project.project_id, "big.txt", big)
    with pytest.raises(UnsupportedFileError):
        container.ingestion.add_document(project.project_id, "empty.txt", b"")


def test_corrupt_file_is_stored_as_failed_not_crashing(container):
    from backend.models.domain import Project

    project = container.tables.projects.put(Project(name="p"))
    art = container.ingestion.add_document(project.project_id, "broken.tf", b'resource "aws_x" "y" {')
    assert art.parse_status == "FAILED" and art.parse_error


def test_encrypted_pdf_is_rejected():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(200, 200)
    w.encrypt("secret")
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(ParseError, match="encrypted"):
        parse_bytes(buf.getvalue(), "locked.pdf", "pdf", LIMITS)


# ----------------------------------------------------------------- text / office / pdf
def test_markdown_sections_have_line_ranges():
    doc = parse_bytes(b"# Intro\nhello\n\n## Data\nPostgres here\nmore\n", "a.md", "markdown", LIMITS)
    heads = [(s.heading, s.line_start, s.line_end) for s in doc.sections]
    assert heads[0][0] == "Intro" and heads[1][0] == "Data"
    assert heads[1][1] == 4 and heads[1][2] == 6


def test_plain_text_rejects_binary():
    with pytest.raises(ParseError):
        parse_bytes(b"\x00\x01\x02binary", "x.txt", "text", LIMITS)


def test_pdf_text_extraction_and_pages():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(72, 750, "Architecture overview: the billing service stores invoices in a PostgreSQL database.")
    c.showPage()
    c.drawString(72, 750, "Disaster recovery: nightly backups are copied to a second region for resilience.")
    c.save()
    doc = parse_bytes(buf.getvalue(), "design.pdf", "pdf", LIMITS)
    assert doc.page_count == 2 and not doc.requires_vision
    assert "PostgreSQL" in doc.sections[0].text and doc.sections[1].page == 2


def test_pdf_without_text_layer_is_flagged_for_vision():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(300, 300)
    buf = io.BytesIO()
    w.write(buf)
    doc = parse_bytes(buf.getvalue(), "scan.pdf", "pdf", LIMITS)
    assert doc.requires_vision and any("OCR is not built in" in w for w in doc.warnings)


def test_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Overview", level=1)
    d.add_paragraph("The API gateway fronts three services.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "RTO"
    t.rows[0].cells[1].text = "30 minutes"
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_bytes(buf.getvalue(), "d.docx", "docx", LIMITS)
    assert doc.sections[0].heading == "Overview"
    assert "RTO | 30 minutes" in doc.text


def test_pptx_slides_and_notes():
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Deployment"
    s.placeholders[1].text = "Kubernetes with two replicas"
    s.notes_slide.notes_text_frame.text = "Mention failover"
    buf = io.BytesIO()
    p.save(buf)
    doc = parse_bytes(buf.getvalue(), "d.pptx", "pptx", LIMITS)
    assert doc.sections[0].slide == 1 and "Kubernetes" in doc.sections[0].text and "Mention failover" in doc.text


def test_zip_bomb_guard():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (3 * 1024 * 1024))
    tiny = ParserLimits(max_zip_uncompressed=1024 * 1024)
    with pytest.raises(ParseError, match="decompression"):
        parse_bytes(buf.getvalue(), "bomb.docx", "docx", tiny)


def test_raster_image_is_not_interpreted():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "white").save(buf, format="PNG")
    doc = parse_bytes(buf.getvalue(), "diagram.png", "image", LIMITS)
    assert doc.requires_vision and not doc.nodes and doc.metadata["width"] == 40


# ----------------------------------------------------------------- diagrams
DRAWIO = """<mxfile><diagram><mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>
<mxCell id="a" value="Web App" vertex="1" parent="1"/><mxCell id="b" value="Orders &lt;b&gt;DB&lt;/b&gt;" vertex="1" parent="1"/>
<mxCell id="e1" value="SQL" edge="1" parent="1" source="a" target="b"/></root></mxGraphModel></diagram></mxfile>"""


def test_drawio_plain_and_compressed():
    doc = parse_bytes(DRAWIO.encode(), "d.drawio", "drawio", LIMITS)
    assert {n.label for n in doc.nodes} == {"Web App", "Orders DB"}
    assert [(e.source, e.target, e.label) for e in doc.edges] == [("a", "b", "SQL")]
    inner = "<mxGraphModel><root><mxCell id='0'/><mxCell id='1' parent='0'/><mxCell id='x' value='Cache' vertex='1' parent='1'/></root></mxGraphModel>"
    comp = zlib.compressobj(wbits=-15)
    payload = base64.b64encode(comp.compress(urllib.parse.quote(inner).encode()) + comp.flush()).decode()
    doc2 = parse_bytes(f"<mxfile><diagram>{payload}</diagram></mxfile>".encode(), "c.drawio", "drawio", LIMITS)
    assert [n.label for n in doc2.nodes] == ["Cache"]


def test_xml_entity_bomb_is_rejected():
    bomb = b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]><mxfile><diagram>&b;</diagram></mxfile>'
    with pytest.raises(ParseError):
        parse_bytes(bomb, "bomb.drawio", "drawio", LIMITS)


def test_mermaid_flowchart_nodes_edges_and_subgraphs():
    src = b"""graph LR
    U([User]) -->|HTTPS| LB[Load Balancer]
    subgraph vpc[Private VPC]
      LB --> App[App Service]
      App -->|SQL| DB[(Orders DB)]
    end
    """
    doc = parse_bytes(src, "a.mmd", "mermaid", LIMITS)
    labels = {n.node_id: n.label for n in doc.nodes}
    assert labels["DB"] == "Orders DB" and labels["LB"] == "Load Balancer"
    assert ("LB", "App") in {(e.source, e.target) for e in doc.edges}
    assert next(n for n in doc.nodes if n.node_id == "App").group == "Private VPC"


def test_plantuml_components_and_arrows():
    src = b"""@startuml
node "Web" as web
database "Orders DB" as db
web --> db : SQL
@enduml"""
    doc = parse_bytes(src, "a.puml", "plantuml", LIMITS)
    assert {n.label for n in doc.nodes} == {"Web", "Orders DB"}
    assert doc.edges[0].label == "SQL"


def test_svg_labels_only_with_limitation_warning():
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><text x="1" y="1">API Gateway</text><text>Database</text></svg>'
    doc = parse_bytes(svg, "a.svg", "svg", LIMITS)
    assert {n.label for n in doc.nodes} == {"API Gateway", "Database"} and not doc.edges
    assert any("connectors" in w for w in doc.warnings)


# ----------------------------------------------------------------- IaC
def test_terraform_resources_have_lines_and_clean_values():
    tf = b'''resource "aws_db_instance" "main" {
  engine = "postgres"
  publicly_accessible = true
  tags = { Name = "orders" }
}
'''
    doc = parse_bytes(tf, "main.tf", "terraform", LIMITS)
    r = doc.resources[0]
    assert r.address == "aws_db_instance.main" and r.line == 1
    assert r.properties["engine"] == "postgres" and r.properties["publicly_accessible"] is True


def test_terraform_heredoc_with_crlf_line_endings():
    tf = b'resource "aws_iam_role" "app" {\r\n  assume_role_policy = <<EOF\r\n{"Version": "2012-10-17"}\r\nEOF\r\n}\r\n'
    doc = parse_bytes(tf, "main.tf", "terraform", LIMITS)
    assert doc.resources[0].address == "aws_iam_role.app"


def test_cloudformation_short_form_tags_and_kubernetes_multidoc():
    cfn = b"""AWSTemplateFormatVersion: '2010-09-09'
Resources:
  Q:
    Type: AWS::SQS::Queue
  F:
    Type: AWS::Lambda::Function
    Properties:
      Role: !GetAtt R.Arn
      Env: !Ref Q
"""
    doc = parse_bytes(cfn, "t.yaml", "structured", LIMITS)
    assert doc.kind == "cloudformation" and {r.address for r in doc.resources} == {"Q", "F"}
    assert doc.resources[1].properties["Env"] == {"Ref": "Q"}
    k8s = b"apiVersion: apps/v1\nkind: Deployment\nmetadata: {name: a}\nspec: {replicas: 1}\n---\napiVersion: v1\nkind: Service\nmetadata: {name: a}\n"
    d2 = parse_bytes(k8s, "k.yaml", "structured", LIMITS)
    assert d2.kind == "kubernetes" and [r.resource_type for r in d2.resources] == ["Deployment", "Service"]
    assert d2.resources[1].line == 6


def test_yaml_alias_expansion_attack_rejected():
    text = "a: &a [x]\n" + "\n".join(f"k{i}: *a" for i in range(300))
    with pytest.raises(ParseError, match="alias"):
        parse_bytes(text.encode(), "bomb.yaml", "structured", LIMITS)


def test_openapi_and_json_documents():
    spec = json.dumps({"openapi": "3.0.0", "info": {"title": "T", "version": "1"},
                       "paths": {"/a": {"get": {"summary": "x"}}}}).encode()
    doc = parse_bytes(spec, "api.json", "structured", LIMITS)
    assert doc.kind == "openapi" and "NO security requirement" in doc.sections[0].text
